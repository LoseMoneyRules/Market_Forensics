from __future__ import annotations

from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

from .report_charts import (
    NAVY, INK, MUTED, LINE, PRIMARY, PRIMARY_SOFT, POSITIVE, NEGATIVE, CAUTION, MARKET,
    chart_bundle,
)


def _n(value: Any) -> float | None:
    try:
        return None if value in (None, "") else float(value)
    except (TypeError, ValueError, ArithmeticError):
        return None


def _money(value: Any) -> str:
    v = _n(value)
    if v is None:
        return "-"
    sign = "-" if v < 0 else ""
    magnitude = abs(v)
    if magnitude >= 1_000_000_000:
        return sign + "$" + f"{magnitude/1_000_000_000:,.1f}B"
    if magnitude >= 1_000_000:
        return sign + "$" + f"{magnitude/1_000_000:,.1f}M"
    if magnitude >= 1_000:
        return sign + "$" + f"{magnitude/1_000:,.1f}K"
    return sign + "$" + f"{magnitude:,.2f}"


def _pct(value: Any, *, signed: bool = True) -> str:
    v = _n(value)
    if v is None:
        return "-"
    return (f"{v:+.1f}%" if signed else f"{v:.1f}%")


def _num(value: Any, digits: int = 1, suffix: str = "") -> str:
    v = _n(value)
    return "-" if v is None else f"{v:,.{digits}f}{suffix}"


def _txt(value: Any, limit: int | None = None) -> str:
    text = " ".join(str(value or "").split())
    if limit and len(text) > limit:
        return text[: max(0, limit - 3)].rstrip() + "..."
    return text or "-"


def _canonicalize_input(data: dict[str, Any]) -> dict[str, Any]:
    """Accept the pre-0.2.14 flat renderer payload without creating a second report engine."""
    if isinstance(data.get("identity"), dict) and isinstance(data.get("valuation"), dict):
        return data
    source = dict(data or {})
    lenses_list = list(source.get("decision_lenses") or [])
    lens_map = {}
    for row in lenses_list:
        if not isinstance(row, dict):
            continue
        key = str(row.get("key") or row.get("label") or "").strip().lower().replace(" ", "_")
        if key:
            lens_map[key] = row.get("state")
    validation_state = source.get("validation_state") or "NOT RUN"
    validation_available = str(validation_state).upper() not in {"", "NOT RUN", "NOT_RUN", "UNAVAILABLE"}
    return {
        "contract_version": "legacy-adapter",
        "mode": "executive" if str(source.get("mode") or "").lower() == "executive" else "full",
        "branding": dict(source.get("branding") or {}),
        "identity": {
            "ticker": source.get("ticker") or "UNKNOWN",
            "company": source.get("company") or "Unknown company",
            "sector": source.get("sector") or "",
            "industry": source.get("industry") or "",
            "market_provider": source.get("market_provider") or "",
            "market_as_of": source.get("market_as_of") or "",
        },
        "conclusion": source.get("action") or "DATA REVIEW",
        "decision_lenses": {
            "value": lens_map.get("value") or source.get("stance") or "UNVERIFIED",
            "expectations": lens_map.get("expectations") or "UNAVAILABLE",
            "variant": lens_map.get("variant") or "UNPROVEN",
            "path": lens_map.get("path") or "UNCLEAR",
            "model_confidence": lens_map.get("model_confidence") or source.get("confidence") or "UNVALIDATED",
            "thesis_control": lens_map.get("thesis_control") or "UNRESOLVED",
            "business": lens_map.get("business") or "UNPROVEN",
            "rows": lenses_list,
        },
        "valuation": {
            "current_price": source.get("market_price"),
            "bear": source.get("bear"), "base": source.get("base"), "bull": source.get("bull"),
            "expected_value": source.get("expected_value"), "base_gap_pct": source.get("base_gap_pct"),
            "quality": source.get("valuation_quality") or "DATA_WARNING",
            "base_quality": source.get("valuation_quality") or "DATA_WARNING",
            "decision_grade": False, "provisional": True, "warnings": [],
            "share_basis": {"shares": None, "source": "UNRESOLVED", "verified": False, "note": ""},
            "weights": {}, "horizon_years": 5, "calibration": {}, "scenarios": [], "engine_version": "",
        },
        "thesis": {
            "thesis": source.get("thesis") or "",
            "counter_evidence": source.get("counter_evidence") or "",
            "market_view": source.get("variant_market") or "",
            "our_view": source.get("variant_us") or "",
            "variant_evidence": source.get("variant_evidence") or "",
            "what_must_be_true": [], "what_proves_wrong": [],
        },
        "evidence": {
            "for": list(source.get("supporting") or []), "against": list(source.get("opposing") or []),
            "score": source.get("score"), "warnings": list(source.get("warnings") or []),
            "blockers": list(source.get("blockers") or []),
        },
        "business": {"summary": source.get("business") or ""},
        "fundamentals": {
            "current": dict(source.get("current_fundamentals") or {}),
            "history": list(source.get("fundamentals_history") or []),
            "summary": source.get("numbers") or "",
        },
        "expectations": {
            "summary": source.get("expectations_summary") or "",
            "implied": dict(source.get("implied_expectations") or {}),
            "rows": list(source.get("expectations") or []),
        },
        "flows": {"summary": source.get("flows_summary") or "", "periods": []},
        "management": {
            "summary": source.get("management_summary") or "", "engine": {}, "accountability": [],
            "promises": list(source.get("management_promises") or []), "assessments": list(source.get("management") or []),
        },
        "catalysts": [{
            "event": row.get("event") or row.get("title"), "timing": row.get("timing") or row.get("date"),
            "direction": row.get("direction"), "status": row.get("status"), "type": row.get("type"),
            "evidence": row.get("evidence"),
        } for row in (source.get("catalysts") or []) if isinstance(row, dict)],
        "bear_case": [{
            "risk": row.get("risk") or row.get("title"), "severity": row.get("severity"),
            "probability": row.get("probability"), "invalidates": row.get("invalidates"),
            "status": row.get("status"), "evidence": row.get("evidence"),
        } for row in (source.get("bear_items") or []) if isinstance(row, dict)],
        "tape": {
            "summary": source.get("tape_summary") or "", "metrics": dict(source.get("tape_metrics") or {}),
            "market": [], "daily_market": [], "short_interest": [], "short_volume": [],
            "institutional_flow": [], "tape_daily": [], "ats": [], "what_changed": "",
            "what_would_change_regime": "",
        },
        "monitoring": {
            "thesis_invalidation": source.get("risk_summary") or "", "locked_at": "",
            "summary": source.get("risk_summary") or "", "rules": [],
        },
        "validation": {
            "state": validation_state, "available": validation_available, "status": validation_state,
            "sample_count": 0, "reliability": None, "lookback_years": None, "history_span": "",
            "valuation_accuracy": None, "direction_accuracy": None, "range_coverage": None,
            "assumption_accuracy": None, "summary": {}, "samples": [],
        },
        "sources": list(source.get("sources") or []), "price_history": [],
        "readiness": {"done": 0, "total": 0, "ready_to_validate": bool(source.get("ready_to_validate"))},
        "triangulation": dict(source.get("triangulation") or {}),
        "data_contract": {"provider_refresh_started": False, "heavy_analytics_started": False},
    }


def _tone(value: str) -> str:
    t = str(value or "").upper()
    if any(x in t for x in ("LONG", "READY", "ATTRACTIVE", "SUPPORTIVE", "VALIDATED", "MET", "GOOD", "STRONG", "POSITIVE")):
        return POSITIVE
    if any(x in t for x in ("SHORT", "FAIL", "MISS", "HOSTILE", "NEGATIVE", "BREACH")):
        return NEGATIVE
    if any(x in t for x in ("WATCH", "REVIEW", "INCOMPLETE", "PROVISIONAL", "WAIT", "CAUTION", "PENDING")):
        return CAUTION
    return PRIMARY


def _flow_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    """Return the newest stored signed financial-flow ledger.

    FinancialFlow persists canonical payloads as edges + signed_exceptions.
    Do not reconstruct an accounting bridge in the report and do not turn
    negative exceptions positive for presentation.
    """
    periods = list((data.get("flows") or {}).get("periods") or [])
    for period in periods:
        rows: list[dict[str, Any]] = []
        for key, flow_label in (("income_statement", "Income statement"), ("cash_flow", "Cash flow")):
            payload = dict(period.get(key) or {})
            for edge in list(payload.get("edges") or []) + list(payload.get("signed_exceptions") or []):
                if not isinstance(edge, dict):
                    continue
                signed = _n(edge.get("signed_value"))
                if signed is None:
                    continue
                source = _txt(edge.get("source"))
                target = _txt(edge.get("target"))
                source_field = str(edge.get("source_field") or "")
                kind = str(edge.get("kind") or "FLOW").upper()
                derived = kind in {"BRIDGE", "WARNING"} or source_field.startswith("derived_")
                rows.append({
                    "period": period.get("period"),
                    "flow": flow_label,
                    "label": edge.get("label") or source_field or "Flow",
                    "value": signed,
                    "route": f"{source} -> {target}",
                    "basis": "Derived bridge" if derived else "Reported / normalized",
                })
        if rows:
            return rows
    return []


def render_pdf_v2(data: dict[str, Any], *, logo_stream: BytesIO | None = None) -> BytesIO:
    data = _canonicalize_input(data)
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image, KeepTogether,
    )

    out = BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=LETTER, leftMargin=.48*inch, rightMargin=.48*inch,
        topMargin=.40*inch, bottomMargin=.55*inch,
        title=f"{(data.get('identity') or {}).get('ticker','')} Market Forensics Research",
        author=(data.get("branding") or {}).get("prepared_by") or "Market Forensics",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="MFBrand", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=8.5, leading=10, textColor=colors.HexColor(PRIMARY), spaceAfter=2, letterSpacing=.6, keepWithNext=1))
    styles.add(ParagraphStyle(name="MFTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=21, leading=24, textColor=colors.HexColor(NAVY), alignment=TA_LEFT, spaceAfter=2))
    styles.add(ParagraphStyle(name="MFMeta", parent=styles["BodyText"], fontSize=8.8, leading=11.5, textColor=colors.HexColor(MUTED), spaceAfter=5))
    styles.add(ParagraphStyle(name="MFHeroLabel", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=8.6, leading=10, textColor=colors.HexColor(MUTED), letterSpacing=.5))
    styles.add(ParagraphStyle(name="MFHero", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=18, leading=21, textColor=colors.HexColor(NAVY)))
    styles.add(ParagraphStyle(name="MFHeroTarget", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=20, leading=22, textColor=colors.HexColor(PRIMARY), alignment=TA_RIGHT))
    styles.add(ParagraphStyle(name="MFH1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=colors.HexColor(NAVY), spaceBefore=10, spaceAfter=5, keepWithNext=1))
    styles.add(ParagraphStyle(name="MFH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11.5, leading=14, textColor=colors.HexColor(PRIMARY), spaceBefore=7, spaceAfter=3, keepWithNext=1))
    styles.add(ParagraphStyle(name="MFBody", parent=styles["BodyText"], fontSize=9.6, leading=13.2, textColor=colors.HexColor(INK), spaceAfter=4))
    styles.add(ParagraphStyle(name="MFSmall", parent=styles["BodyText"], fontSize=8.7, leading=11.2, textColor=colors.HexColor(MUTED), spaceAfter=2))
    styles.add(ParagraphStyle(name="MFKpiL", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=7.4, leading=9, textColor=colors.HexColor(MUTED), alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="MFKpiV", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=10.4, leading=12.2, textColor=colors.HexColor(NAVY), alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="MFCell", parent=styles["BodyText"], fontSize=8.8, leading=11.2, textColor=colors.HexColor(INK)))
    styles.add(ParagraphStyle(name="MFCellSmall", parent=styles["BodyText"], fontSize=8.1, leading=10.2, textColor=colors.HexColor(INK)))
    styles.add(ParagraphStyle(name="MFSource", parent=styles["BodyText"], fontSize=6.6, leading=8.1, textColor=colors.HexColor(MUTED)))
    styles.add(ParagraphStyle(name="MFTableHead", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=8.0, leading=9.7, textColor=colors.HexColor(MUTED)))
    styles.add(ParagraphStyle(name="MFCenter", parent=styles["MFCell"], alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="MFRight", parent=styles["MFCell"], alignment=TA_RIGHT))

    brand = data.get("branding") or {}
    ident = data.get("identity") or {}
    valuation = data.get("valuation") or {}
    lenses = data.get("decision_lenses") or {}
    evidence = data.get("evidence") or {}
    charts = chart_bundle(data)
    story = []

    def P(text: Any, style: str = "MFBody"):
        return Paragraph(escape(_txt(text)).replace("\n", "<br/>"), styles[style])

    def rich(text: str, style: str = "MFBody"):
        return Paragraph(text, styles[style])

    def rule_table(rows, widths=None, *, header=True, font_style="MFCell", repeat=1):
        wrapped = []
        for r_idx, row in enumerate(rows):
            style_name = "MFTableHead" if header and r_idx == 0 else font_style
            wrapped.append([Paragraph(escape(_txt(cell)), styles[style_name]) for cell in row])
        t = Table(wrapped, colWidths=widths, repeatRows=repeat if header else 0, hAlign="LEFT")
        cmds = [
            ("VALIGN", (0,0), (-1,-1), "TOP"),
            ("LINEBELOW", (0,0), (-1,-1), .28, colors.HexColor("#e4e9ed")),
            ("LEFTPADDING",(0,0),(-1,-1),5), ("RIGHTPADDING",(0,0),(-1,-1),5),
            ("TOPPADDING",(0,0),(-1,-1),4), ("BOTTOMPADDING",(0,0),(-1,-1),4),
        ]
        if header:
            cmds += [("BACKGROUND",(0,0),(-1,0),colors.HexColor("#f2f5f7")),("LINEBELOW",(0,0),(-1,0),.6,colors.HexColor(LINE))]
        t.setStyle(TableStyle(cmds))
        return t

    def section(title: str, eyebrow: str | None = None):
        # Both styles use keepWithNext: eyebrow -> heading -> first section
        # content. Keeping them as separate flowables lets ReportLab move the
        # whole chain instead of orphaning a label at the page bottom.
        if eyebrow:
            story.append(Paragraph(escape(eyebrow.upper()), styles["MFBrand"]))
        story.append(Paragraph(escape(title), styles["MFH1"]))

    def kpi_strip(items: list[tuple[str,str,str]], cols: int = 5):
        items = list(items)
        rows = []
        for offset in range(0, len(items), cols):
            chunk = items[offset:offset+cols]
            while len(chunk) < cols:
                chunk.append(("", "", NAVY))
            rows.append([Paragraph(escape(lbl.upper()), styles["MFKpiL"]) for lbl,_,_ in chunk])
            vals = []
            for _, val, tone in chunk:
                # Preserve the restrained semantic state colors used by the Web UI.
                para = Paragraph(
                    f'<font color="{escape(str(tone or NAVY))}">{escape(val or "-")}</font>',
                    styles["MFKpiV"],
                )
                vals.append(para)
            rows.append(vals)
        widths = [6.55*inch/cols]*cols
        t = Table(rows, colWidths=widths)
        cmds = [
            ("BOX",(0,0),(-1,-1),.45,colors.HexColor(LINE)),
            ("INNERGRID",(0,0),(-1,-1),.25,colors.HexColor(LINE)),
            ("BACKGROUND",(0,0),(-1,-1),colors.white),
            ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
            ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4),
            ("LEFTPADDING",(0,0),(-1,-1),2),("RIGHTPADDING",(0,0),(-1,-1),2),
        ]
        for r_idx in range(0, len(rows), 2):
            cmds.append(("BACKGROUND",(0,r_idx),(-1,r_idx),colors.HexColor("#f2f5f7")))
        t.setStyle(TableStyle(cmds))
        story.extend([t, Spacer(1,5)])

    def image(key: str, width: float = 6.55, max_height: float | None = None):
        stream = charts.get(key)
        if not stream:
            return
        stream.seek(0)
        from PIL import Image as PILImage
        with PILImage.open(stream) as im:
            aspect = im.height / im.width
        stream.seek(0)
        h = width * aspect
        if max_height and h > max_height:
            h = max_height
            width = h / aspect
        story.extend([Image(stream, width=width*inch, height=h*inch), Spacer(1,5)])

    # Header / identity.
    if logo_stream is not None:
        try:
            logo_stream.seek(0)
            story.append(Image(logo_stream, width=.34*inch, height=.34*inch))
        except Exception:
            pass
    story.append(Paragraph(escape(str(brand.get("title") or "MARKET FORENSICS").upper()), styles["MFBrand"]))
    story.append(Paragraph(escape(f"{ident.get('ticker','')}  |  {ident.get('company','')}"), styles["MFTitle"]))
    meta = "  |  ".join(x for x in [ident.get("sector"), ident.get("industry"), ident.get("market_provider"), ident.get("market_as_of")] if x)
    story.append(Paragraph(escape(meta or "Private CONTROL research"), styles["MFMeta"]))

    # First decision information: conclusion hero.
    conclusion = str(data.get("conclusion") or "DATA REVIEW")
    base = _money(valuation.get("base"))
    gap = _pct(valuation.get("base_gap_pct"))
    hero_left = [
        Paragraph("RESEARCH CONCLUSION", styles["MFHeroLabel"]),
        Paragraph(f'<font color="{_tone(conclusion)}">{escape(conclusion)}</font>', styles["MFHero"]),
        Paragraph(escape(f"Value {lenses.get('value','-')}  |  Path {lenses.get('path','-')}  |  Confidence {lenses.get('model_confidence','-')}"), styles["MFSmall"]),
    ]
    hero_right = [
        Paragraph("BASE TARGET", styles["MFHeroLabel"]),
        Paragraph(escape(base), styles["MFHeroTarget"]),
        Paragraph(escape(f"{gap} vs market  |  {_txt(valuation.get('base_quality')).replace('_',' ')}"), styles["MFSmall"]),
    ]
    hero = Table([[hero_left, hero_right]], colWidths=[4.15*inch, 2.4*inch])
    hero.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#f7fafc")),
        ("BOX",(0,0),(-1,-1),.7,colors.HexColor("#bdd0df")),
        ("LINEBEFORE",(1,0),(1,0),.7,colors.HexColor("#bdd0df")),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("LEFTPADDING",(0,0),(-1,-1),10),("RIGHTPADDING",(0,0),(-1,-1),10),
        ("TOPPADDING",(0,0),(-1,-1),9),("BOTTOMPADDING",(0,0),(-1,-1),9),
    ]))
    story.extend([hero, Spacer(1,6)])

    # The decision range is intentionally visual and immediate: conclusion -> valuation map
    # -> compact decision diagnostics. This mirrors the first-read order used in sell-side/buy-side notes.
    image("valuation_map", 6.55, 1.58)
    kpi_strip([
        ("Current price", _money(valuation.get("current_price")), MARKET),
        ("Bear", _money(valuation.get("bear")), NEGATIVE),
        ("Base", _money(valuation.get("base")), PRIMARY),
        ("Bull", _money(valuation.get("bull")), POSITIVE),
        ("Base gap", _pct(valuation.get("base_gap_pct")), PRIMARY),
        ("Valuation quality", _txt(valuation.get("base_quality")).replace("_"," "), _tone(str(valuation.get("base_quality")))),
        ("Value lens", _txt(lenses.get("value")), _tone(str(lenses.get("value")))),
        ("Expectations", _txt(lenses.get("expectations")), _tone(str(lenses.get("expectations")))),
        ("Variant", _txt(lenses.get("variant")), _tone(str(lenses.get("variant")))),
        ("Path", _txt(lenses.get("path")), _tone(str(lenses.get("path")))),
        ("Model confidence", _txt(lenses.get("model_confidence")), _tone(str(lenses.get("model_confidence")))),
        ("Thesis control", _txt(lenses.get("thesis_control")), _tone(str(lenses.get("thesis_control")))),
    ], 6)

    # Executive decision content.
    thesis = data.get("thesis") or {}
    executive_quad = [
        [
            rich("<b>MARKET VIEW</b><br/>" + escape(_txt(thesis.get("market_view"), 360)), "MFBody"),
            rich("<b>OUR VIEW / VARIANT</b><br/>" + escape(_txt(thesis.get("our_view"), 360)), "MFBody"),
        ],
        [
            rich("<b>WHAT MUST BE TRUE</b><br/>" + "<br/>".join("- "+escape(_txt(x.get("text"), 160)) for x in (thesis.get("what_must_be_true") or [])[:3]) if thesis.get("what_must_be_true") else "<b>WHAT MUST BE TRUE</b><br/>-", "MFBody"),
            rich("<b>WHAT WOULD PROVE US WRONG</b><br/>" + "<br/>".join("- "+escape(_txt(x.get("text"), 160)) for x in (thesis.get("what_proves_wrong") or [])[:3]) if thesis.get("what_proves_wrong") else "<b>WHAT WOULD PROVE US WRONG</b><br/>-", "MFBody"),
        ],
    ]
    quad = Table(executive_quad, colWidths=[3.25*inch,3.25*inch])
    quad.setStyle(TableStyle([
        ("BOX",(0,0),(-1,-1),.4,colors.HexColor(LINE)),("INNERGRID",(0,0),(-1,-1),.3,colors.HexColor(LINE)),
        ("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#fbfcfd")),("VALIGN",(0,0),(-1,-1),"TOP"),
        ("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),7),
        ("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6),
    ]))
    story.extend([quad, Spacer(1,5)])

    def evidence_html(rows):
        rows = list(rows or [])[:4]
        if not rows:
            return "-"
        return "<br/>".join("<b>"+escape(_txt(r.get("label") or "Evidence",70))+"</b> - "+escape(_txt(r.get("detail"),180)) for r in rows)
    ev = Table([
        [rich("<b>FOR</b><br/>"+evidence_html(evidence.get("for")), "MFBody"),
         rich("<b>AGAINST</b><br/>"+evidence_html(evidence.get("against")), "MFBody")]
    ], colWidths=[3.25*inch,3.25*inch])
    ev.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(0,0),colors.HexColor("#edf7f1")),("BACKGROUND",(1,0),(1,0),colors.HexColor("#fbefef")),
        ("BOX",(0,0),(-1,-1),.35,colors.HexColor(LINE)),("VALIGN",(0,0),(-1,-1),"TOP"),
        ("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),7),
        ("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6),
    ]))
    story.extend([ev, Spacer(1,4)])

    company_quality = data.get("company_quality") or valuation.get("company_quality") or {}
    if company_quality:
        alarms = list(company_quality.get("alarm_bells") or [])
        strengths = list(company_quality.get("strengths") or [])
        quality_rows = [
            ["Company quality", _txt(company_quality.get("state") or "INSUFFICIENT EVIDENCE") + "  |  " + _txt(company_quality.get("headline"), 250)],
            ["Alarm bells", " | ".join(_txt(x.get("detail"), 150) for x in alarms[:2]) or "No material automatic alarm bell."],
            ["Strengths", " | ".join(_txt(x.get("detail"), 150) for x in strengths[:2]) or "No automatic strength clears the threshold yet."],
        ]
        story.append(rule_table(quality_rows, widths=[1.0*inch,5.55*inch], header=False))
        story.append(Spacer(1,4))

    tape = data.get("tape") or {}
    tm = tape.get("metrics") or {}
    catalysts = data.get("catalysts") or []
    next_cat = catalysts[0] if catalysts else {}
    invalidation = (data.get("monitoring") or {}).get("thesis_invalidation")
    exec_rows = [
        ["Tape", "  |  ".join(filter(None, [
            f"Regime {_txt(tm.get('regime'))}", f"Rank {_txt(tm.get('rank'))}",
            f"Confidence {_txt(tm.get('confidence'))}", f"Net Tape {_num(tm.get('net_tape'),0)}"
        ]))],
        ["Next catalyst", "  |  ".join(filter(None, [_txt(next_cat.get("event")), _txt(next_cat.get("timing")), _txt(next_cat.get("status"))]))],
        ["Invalidation", _txt(invalidation, 260)],
    ]
    story.append(rule_table(exec_rows, widths=[1.0*inch,5.55*inch], header=False))
    score = evidence.get("score")
    if score is not None:
        story.append(Paragraph(escape(f"Evidence diagnostic {float(score):+.2f} - secondary diagnostic only; Decision Lenses remain canonical."), styles["MFSmall"]))

    if data.get("mode") == "executive":
        def footer(canvas, _doc):
            canvas.saveState()
            canvas.setStrokeColor(colors.HexColor(LINE)); canvas.setLineWidth(.35)
            canvas.line(.48*inch,.38*inch,8.02*inch,.38*inch)
            canvas.setFillColor(colors.HexColor(MUTED)); canvas.setFont("Helvetica",7.4)
            canvas.drawString(.48*inch,.22*inch,f"{brand.get('footer') or 'Lose Money Rules'} | Market Forensics | {ident.get('ticker','')} | Executive")
            canvas.drawRightString(8.02*inch,.22*inch,f"Page {canvas.getPageNumber()}")
            canvas.restoreState()
        doc.build(story, onFirstPage=footer, onLaterPages=footer)
        out.seek(0)
        return out

    # Full report begins after executive read.
    story.append(PageBreak())

    section("Valuation", "Decision economics")
    story.append(Paragraph(
        escape("PRICE -> BEAR / BASE / BULL -> GAP. Detail follows only after the decision range."),
        styles["MFSmall"],
    ))
    kpi_strip([
        ("Current price", _money(valuation.get("current_price")), MARKET),
        ("Bear", _money(valuation.get("bear")), NEGATIVE),
        ("Base", _money(valuation.get("base")), PRIMARY),
        ("Bull", _money(valuation.get("bull")), POSITIVE),
        ("Expected value", _money(valuation.get("expected_value")), NAVY),
    ], 5)
    image("price_context", 6.55, 2.1)
    rows = [["Case","Probability","Target","Quality","P/E","EV/Sales","FCF Yield","DCF"]]
    for row in valuation.get("scenarios") or []:
        methods = {m.get("key"): m for m in row.get("methods") or []}
        rows.append([
            row.get("name"), _pct((_n(row.get("probability")) or 0)*100, signed=False),
            _money(row.get("target")), _txt(row.get("quality")).replace("_"," "),
            _money((methods.get("pe") or {}).get("value")),
            _money((methods.get("ev_sales") or {}).get("value")),
            _money((methods.get("fcf_yield") or {}).get("value")),
            _money(row.get("dcf")),
        ])
    story.append(rule_table(rows, widths=[.55*inch,.70*inch,.75*inch,1.18*inch,.78*inch,.78*inch,.78*inch,.78*inch], font_style="MFCellSmall"))
    weight_rows = [["Case","P/E weight","EV/Sales weight","FCF Yield weight"]]
    for row in valuation.get("scenarios") or []:
        methods = {m.get("key"): m for m in row.get("methods") or []}
        def method_weight(key):
            weight = _n((methods.get(key) or {}).get("weight"))
            return _pct(weight * 100, signed=False) if weight is not None else "-"
        weight_rows.append([
            row.get("name"), method_weight("pe"), method_weight("ev_sales"), method_weight("fcf_yield")
        ])
    if len(weight_rows) > 1:
        story.append(Paragraph("EFFECTIVE METHOD WEIGHTS", styles["MFBrand"]))
        story.append(rule_table(weight_rows, widths=[1.0*inch,1.8*inch,1.8*inch,1.95*inch], font_style="MFCellSmall"))
    share = valuation.get("share_basis") or {}
    story.append(Paragraph(
        escape("Share denominator: " + _num(share.get("shares"),1) + " | Source: " + _txt(share.get("source")).replace("_"," ") +
               " | Verified: " + ("YES" if share.get("verified") else "NO") +
               ((" | " + _txt(share.get("note"),160)) if share.get("note") else "")),
        styles["MFSmall"],
    ))
    if valuation.get("provisional"):
        story.append(Paragraph("<b>PROVISIONAL VALUATION.</b> Stored targets remain visible, but this Base is not decision-grade intrinsic evidence until the quality issue is resolved.", styles["MFBody"]))
    for warning in (valuation.get("warnings") or [])[:4]:
        story.append(Paragraph("WATCH - " + escape(_txt(warning, 420)), styles["MFSmall"]))

    policy = valuation.get("valuation_policy") or {}
    ledger = valuation.get("valuation_impact_ledger") or []
    section("Quality → valuation", "Explicit automatic price adjustments")
    kpi_strip([
        ("Company quality", _txt((valuation.get("company_quality") or {}).get("state") or "INSUFFICIENT EVIDENCE"), NAVY),
        ("Risk premium", "+" + str(int(_n(policy.get("risk_premium_bps")) or 0)) + " bps", NAVY),
        ("Growth haircut", "-" + str(int(_n(policy.get("growth_haircut_bps")) or 0)) + " bps", NAVY),
        ("Terminal haircut", "-" + str(int(_n(policy.get("terminal_growth_haircut_bps")) or 0)) + " bps", NAVY),
        ("Bear probability", "+" + str(int(_n(policy.get("bear_probability_shift_pts")) or 0)) + " pts", NAVY),
    ], 5)
    if ledger:
        rows = [["Item","Effect","Impact","Reason"]]
        for item in ledger[:10]:
            rows.append([item.get("item"), _txt(item.get("effect")).replace("_"," "), item.get("impact"), _txt(item.get("reason"),280)])
        story.append(rule_table(rows, widths=[1.35*inch,1.05*inch,1.75*inch,2.4*inch], font_style="MFCellSmall"))
    story.append(Paragraph("Positive company quality receives no automatic premium. Reported operating strength already enters through growth, margins, returns and cash flows.", styles["MFSmall"]))

    section("Thesis / variant", "Why the market may be wrong")
    story.append(quad)
    if thesis.get("variant_evidence"):
        story.append(Paragraph("<b>VARIANT EVIDENCE</b> - "+escape(_txt(thesis.get("variant_evidence"),900)), styles["MFBody"]))

    section("Expectations", "What is priced")
    implied = (data.get("expectations") or {}).get("implied") or {}
    if implied.get("available"):
        rows = [["Driver","Market-implied","Our Base","Read"]]
        for row in implied.get("drivers") or []:
            unit=row.get("unit"); mv=_n(row.get("market_implied")); bv=_n(row.get("base"))
            rows.append([
                row.get("label"),
                _pct(mv*100, signed=False) if unit=="%" and mv is not None else _num(mv,1,"x"),
                _pct(bv*100, signed=False) if unit=="%" and bv is not None else _num(bv,1,"x"),
                row.get("read"),
            ])
        story.append(rule_table(rows, widths=[2.25*inch,1.25*inch,1.25*inch,1.8*inch]))
    exp_rows=(data.get("expectations") or {}).get("rows") or []
    if exp_rows:
        rows=[["Metric","Period","Market","Our view","Confidence"]]
        for row in exp_rows[:12]:
            rows.append([row.get("metric"),row.get("period"),str(row.get("market") if row.get("market") is not None else "-"),str(row.get("ours") if row.get("ours") is not None else "-"),row.get("confidence")])
        story.append(rule_table(rows, widths=[2.1*inch,.85*inch,1.05*inch,1.05*inch,1.5*inch]))
    elif not implied.get("available"):
        story.append(P((data.get("expectations") or {}).get("summary") or "Expectation evidence unavailable."))

    section("Evidence for / against", "Research evidence")
    story.append(ev)
    if evidence.get("warnings") or evidence.get("blockers"):
        for item in (evidence.get("blockers") or [])[:4]:
            story.append(Paragraph("<b>BLOCKER</b> - "+escape(_txt(item,360)), styles["MFSmall"]))
        for item in (evidence.get("warnings") or [])[:4]:
            story.append(Paragraph("<b>WATCH</b> - "+escape(_txt(item,360)), styles["MFSmall"]))

    section("Company quality", "Is the economic engine actually good?")
    cq = data.get("company_quality") or {}
    if cq:
        story.append(Paragraph("<b>"+escape(_txt(cq.get("state")))+"</b> - "+escape(_txt(cq.get("headline"),520)), styles["MFBody"]))
        qrows = [["Dimension","State","Read"]]
        for d in cq.get("dimensions") or []:
            qrows.append([d.get("label"), _txt(d.get("state")).replace("_"," "), _txt(d.get("detail"),300)])
        story.append(rule_table(qrows, widths=[1.45*inch,1.0*inch,4.1*inch], font_style="MFCellSmall"))
        for flag in (cq.get("alarm_bells") or [])[:6]:
            story.append(Paragraph("<b>ALARM · "+escape(_txt(flag.get("severity")))+"</b> - "+escape(_txt(flag.get("detail"),380)), styles["MFSmall"]))
    else:
        story.append(P("Company-quality evidence is not materialized yet."))

    section("Business", "Operating reality")
    story.append(P((data.get("business") or {}).get("summary") or "Business research is not yet documented."))

    section("Fundamentals", "Trend before table")
    ff=(data.get("fundamentals") or {}).get("forensics") or {}
    if ff:
        story.append(Paragraph("<b>"+escape(_txt(ff.get("state") or "DATA REVIEW"))+"</b> - "+escape(_txt(ff.get("headline"),520)), styles["MFBody"]))
        strengths=list(ff.get("strengths") or [])[:4]
        risks=(list(ff.get("red_flags") or []) + list(ff.get("watches") or []))[:5]
        if strengths or risks:
            left="<b>STRENGTHS</b><br/>"+("<br/>".join("- "+escape(_txt(x.get("label"),80))+": "+escape(_txt(x.get("detail"),180)) for x in strengths) if strengths else "-")
            right="<b>RED FLAGS / WATCH</b><br/>"+("<br/>".join("- "+escape(_txt(x.get("label"),80))+": "+escape(_txt(x.get("detail"),180)) for x in risks) if risks else "-")
            ft=Table([[rich(left,"MFBody"),rich(right,"MFBody")]],colWidths=[3.25*inch,3.25*inch])
            ft.setStyle(TableStyle([
                ("BACKGROUND",(0,0),(0,0),colors.HexColor("#edf7f1")),("BACKGROUND",(1,0),(1,0),colors.HexColor("#fbefef")),
                ("BOX",(0,0),(-1,-1),.35,colors.HexColor(LINE)),("VALIGN",(0,0),(-1,-1),"TOP"),
                ("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),7),
                ("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6),
            ]))
            story.extend([ft,Spacer(1,4)])
        for item in (list(ff.get("inconsistencies") or []) + list(ff.get("data_gaps") or []))[:5]:
            story.append(Paragraph("<b>REVIEW</b> - "+escape(_txt(item.get("label"),90))+": "+escape(_txt(item.get("detail"),340)), styles["MFSmall"]))
    image("revenue_profitability", 6.55, 2.75)
    image("cash_conversion", 6.55, 2.3)
    image("working_capital", 6.55, 2.1)
    current=(data.get("fundamentals") or {}).get("current") or {}
    kpi_strip([
        ("Basis", _txt(current.get("period") or current.get("comparison_basis")), NAVY),
        ("Op margin", _pct(current.get("operating_margin_pct"), signed=False), NAVY),
        ("FCF margin", _pct(current.get("fcf_margin_pct"), signed=False), NAVY),
        ("CFO / NI", _num(current.get("cfo_to_net_income"),2,"x"), NAVY),
        ("ROIC", _pct(current.get("roic_pct"), signed=False), NAVY),
        ("Net debt / FCF", _num(current.get("net_debt_to_fcf"),1,"x"), NAVY),
        ("Inventory / Rev", _pct(current.get("inventory_to_revenue_pct"), signed=False), NAVY),
        ("Receivables / Rev", _pct(current.get("receivables_to_revenue_pct"), signed=False), NAVY),
        ("CCC", _num(current.get("ccc"),0,"d"), NAVY),
        ("Share count growth", _pct(current.get("share_count_growth_pct")), NAVY),
    ], 5)
    if current.get("economic_reality_quality"):
        story.append(Paragraph("<b>ECONOMIC REALITY</b> · reported accounting kept intact; financing and operating obligations are classified before scoring.", styles["MFSmall"]))
        kpi_strip([
            ("Econ quality", _txt(current.get("economic_reality_quality")), NAVY),
            ("Reported net debt", _money(current.get("reported_net_debt")), NAVY),
            ("Economic net debt", _money(current.get("economic_net_debt")), NAVY),
            ("Operating leases", _money(current.get("operating_lease_liability")), NAVY),
            ("Lease / liabilities", _pct(current.get("operating_lease_share_of_liabilities_pct"), signed=False), NAVY),
            ("Revenue / leases", _num(current.get("lease_revenue_productivity_x"),2,"x"), NAVY),
            ("Growth capex proxy", _money(current.get("growth_capex_proxy")), NAVY),
            ("Owner-cash proxy", _money(current.get("owner_cash_proxy")), NAVY),
            ("FCF after SBC", _money(current.get("fcf_after_sbc")), NAVY),
            ("Lease-adj ROIC", _pct(current.get("lease_adjusted_roic_pct"), signed=False), NAVY),
        ], 5)
        for flag in (current.get("economic_reality_flags") or [])[:6]:
            story.append(Paragraph("<b>"+escape(_txt(flag.get("code")).replace("_"," "))+"</b> · "+escape(_txt(flag.get("detail"),360)), styles["MFSmall"]))
        if current.get("economic_reality_unresolved"):
            story.append(Paragraph("<b>ACCOUNTING REVIEW</b> · material classification is unresolved; directional scoring is conservative until verified.", styles["MFSmall"]))
    hist=(data.get("fundamentals") or {}).get("history") or []
    if hist:
        rows=[["Period","Revenue growth","Gross M","Op M","FCF M","CFO/NI","ROIC","Inv/Rev","Rec/Rev","CCC"]]
        for r in hist[-8:]:
            rows.append([
                r.get("period"),_pct(r.get("revenue_growth_pct")),_pct(r.get("gross_margin_pct"),signed=False),
                _pct(r.get("operating_margin_pct"),signed=False),_pct(r.get("fcf_margin_pct"),signed=False),
                _num(r.get("cfo_to_net_income"),2,"x"),_pct(r.get("roic_pct"),signed=False),
                _pct(r.get("inventory_to_revenue_pct"),signed=False),_pct(r.get("receivables_to_revenue_pct"),signed=False),
                _num(r.get("ccc"),0,"d"),
            ])
        story.append(rule_table(rows, widths=[.53*inch,.68*inch,.62*inch,.58*inch,.58*inch,.59*inch,.52*inch,.67*inch,.67*inch,.51*inch], font_style="MFCellSmall"))
    summary=(data.get("fundamentals") or {}).get("summary")
    if summary:
        story.append(P(summary))

    section("Financial flows", "Follow the money")
    flow_steps=_flow_rows(data)
    if flow_steps:
        rows=[["Flow","Line item","Signed amount","Route","Basis"]]
        for step in flow_steps[:16]:
            rows.append([
                step.get("flow"), step.get("label"), _money(step.get("value")),
                step.get("route"), step.get("basis"),
            ])
        story.append(rule_table(rows, widths=[.9*inch,1.3*inch,1.0*inch,2.0*inch,1.35*inch], font_style="MFCellSmall"))
        story.append(Paragraph("Signed negatives are retained exactly as stored. The report does not invent balancing values to force a bridge.",styles["MFSmall"]))
    else:
        story.append(P((data.get("flows") or {}).get("summary") or "Materialized financial-flow bridge unavailable."))

    story.append(PageBreak())
    section("Management", "Execution accountability")
    mg=data.get("management") or {}
    eng=mg.get("engine") or {}
    kpi_strip([
        ("Execution", _txt(eng.get("label") or eng.get("execution_label") or "LOW DATA"), _tone(str(eng.get("label")))),
        ("Evidence coverage", _pct(eng.get("coverage_pct"), signed=False), NAVY),
        ("Promises", str(len(mg.get("promises") or [])), NAVY),
        ("Scorable", str(eng.get("scorable") or "-"), NAVY),
        ("Missed", str(eng.get("missed") or "-"), NEGATIVE),
    ],5)
    promises=mg.get("promises") or []
    if promises:
        rows=[["Period","Metric","Promise","Actual","Status","Comparability"]]
        for row in promises[:16]:
            lo=row.get("low");hi=row.get("high");unit=row.get("unit") or ""
            promise=_txt(row.get("target_text"))
            if promise=="-" and lo is not None:
                promise=str(lo) if lo==hi else f"{lo} - {hi}"
                if unit:
                    promise += " "+str(unit)
            rows.append([
                row.get("target_period") or (("FY"+str(row.get("target_year"))) if row.get("target_year") else "UNRESOLVED"),
                row.get("metric"),promise,str(row.get("actual") if row.get("actual") is not None else "-"),
                row.get("status"),row.get("comparability") or row.get("comparability_reason") or "-",
            ])
        story.append(rule_table(rows, widths=[.8*inch,1.35*inch,1.3*inch,.8*inch,.75*inch,1.55*inch], font_style="MFCellSmall"))
    elif mg.get("summary"):
        story.append(P(mg.get("summary")))
    else:
        story.append(P("No materialized management promise ledger is available."))

    section("Catalysts / bear case", "Path and failure modes")
    cats=data.get("catalysts") or []
    if cats:
        rows=[["Event","Timing","Direction","Status"]]
        for row in cats[:12]:
            rows.append([row.get("event"),row.get("timing"),row.get("direction"),row.get("status")])
        story.append(rule_table(rows, widths=[3.45*inch,1.1*inch,.95*inch,1.05*inch]))
    bears=data.get("bear_case") or []
    if bears:
        story.append(Spacer(1,5))
        rows=[["Risk","Severity","Status","Invalidates","Why it matters"]]
        for row in bears[:12]:
            rows.append([row.get("risk"),row.get("severity"),row.get("status"),"YES" if row.get("invalidates") else "NO",_txt(row.get("evidence"),180)])
        story.append(rule_table(rows, widths=[2.0*inch,.75*inch,.75*inch,.70*inch,2.35*inch], font_style="MFCellSmall"))

    section("Tape & positioning", "Market evidence")
    kpi_strip([
        ("Regime", _txt(tm.get("regime") or tm.get("forensic_regime")), _tone(str(tm.get("regime")))),
        ("Rank", _txt(tm.get("rank")), NAVY),
        ("Confidence", _txt(tm.get("confidence")), NAVY),
        ("Large flow", _money(tm.get("net_large")), _tone("POSITIVE" if (_n(tm.get("net_large")) or 0)>=0 else "NEGATIVE")),
        ("Whale flow", _money(tm.get("net_whale")), _tone("POSITIVE" if (_n(tm.get("net_whale")) or 0)>=0 else "NEGATIVE")),
        ("Short pressure", _num(tm.get("bear_pressure"),0), NEGATIVE),
        ("Absorption", _num(tm.get("absorption"),0), POSITIVE),
        ("Net Tape", _num(tm.get("net_tape"),0), PRIMARY),
        ("What changed", _txt(tape.get("what_changed"),60), NAVY),
        ("Regime change", _txt(tape.get("what_would_change_regime"),60), NAVY),
    ],5)
    image("tape_price_flow",6.55,2.5)
    image("tape_pressure",6.55,2.1)
    if tape.get("summary"):
        story.append(P(tape.get("summary")))

    section("Thesis invalidation / monitoring", "Decision control")
    monitoring=data.get("monitoring") or {}
    story.append(Paragraph("<b>THESIS INVALIDATION</b> - "+escape(_txt(monitoring.get("thesis_invalidation"),700)),styles["MFBody"]))
    if monitoring.get("locked_at"):
        story.append(Paragraph(escape("Locked pre-investment basis: "+str(monitoring.get("locked_at"))+". Thresholds are not rewritten retroactively."),styles["MFSmall"]))
    rules=monitoring.get("rules") or []
    if rules:
        rows=[["Metric / rule","Threshold","Direction","Current","State","Last observation","Locked"]]
        for row in rules[:20]:
            threshold=_txt(row.get("threshold_text"))
            if threshold=="-" and row.get("threshold") is not None:
                threshold=_num(row.get("threshold"),2)+((" "+row.get("unit")) if row.get("unit") else "")
            rows.append([
                row.get("name") or row.get("metric"),threshold,row.get("operator"),
                _num(row.get("current_value"),2)+((" "+row.get("unit")) if row.get("unit") and row.get("current_value") is not None else ""),
                row.get("status"),str(row.get("last_observation") or "")[:19],"YES" if row.get("locked_pre_investment") else "NO",
            ])
        story.append(rule_table(rows, widths=[1.55*inch,1.0*inch,.55*inch,.8*inch,.75*inch,1.35*inch,.55*inch], font_style="MFCellSmall"))
    else:
        story.append(P("No active materialized monitoring rules."))

    section("Validation", "Historical calibration")
    val=data.get("validation") or {}
    if not val.get("available"):
        story.append(Paragraph("<b>NOT RUN.</b> No historical validation result is stored. No validation confidence is implied.",styles["MFBody"]))
    else:
        kpi_strip([
            ("State", _txt(val.get("state")), _tone(str(val.get("state")))),
            ("Reliability", _pct(val.get("reliability"), signed=False), NAVY),
            ("Samples", str(val.get("sample_count") or 0), NAVY),
            ("Valuation accuracy", _pct(val.get("valuation_accuracy"), signed=False), NAVY),
            ("Direction accuracy", _pct(val.get("direction_accuracy"), signed=False), NAVY),
            ("Range coverage", _pct(val.get("range_coverage"), signed=False), NAVY),
            ("Assumption accuracy", _pct(val.get("assumption_accuracy"), signed=False), NAVY),
            ("Requested span", (str(val.get("lookback_years") or "-")+"Y"), NAVY),
            ("History span", _txt(val.get("history_span")), NAVY),
            ("Status", _txt(val.get("status")), NAVY),
        ],5)
        image("validation",6.55,2.25)
        summary=val.get("summary") or {}
        insight=summary.get("calibration_insight") or summary.get("bias") or summary.get("message")
        if insight:
            story.append(Paragraph("<b>Calibration insight</b> - "+escape(_txt(insight,600)),styles["MFBody"]))

    section("Sources / audit", "Provenance")
    sources=data.get("sources") or []
    if sources:
        rows=[["Provider","Document / type","Published","Retrieved","Title"]]
        for row in sources[:50]:
            doc_type=" / ".join(x for x in [row.get("document"),row.get("type")] if x)
            rows.append([row.get("provider"),doc_type,str(row.get("published_at") or "")[:10],str(row.get("retrieved_at") or "")[:19],_txt(row.get("title"),140)])
        # Complete provenance, deliberately de-emphasized like a bank research source appendix.
        story.append(rule_table(rows,widths=[.8*inch,1.2*inch,.75*inch,1.0*inch,2.8*inch],font_style="MFSource"))
    else:
        story.append(P("No source provenance rows are stored."))

    contract=data.get("data_contract") or {}
    story.append(Spacer(1,4))
    story.append(Paragraph(escape(
        f"Report data contract {data.get('contract_version','')} | Cache {contract.get('materialized_cache_event') or '-'} | "
        f"Provider refresh started: {'YES' if contract.get('provider_refresh_started') else 'NO'} | "
        f"Heavy analytics started: {'YES' if contract.get('heavy_analytics_started') else 'NO'}"
    ),styles["MFSmall"]))

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor(LINE))
        canvas.setLineWidth(.35)
        canvas.line(.48*inch,.38*inch,8.02*inch,.38*inch)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.setFont("Helvetica",7.4)
        canvas.drawString(.48*inch,.22*inch,f"{brand.get('footer') or 'Lose Money Rules'} | Market Forensics | {ident.get('ticker','')}")
        canvas.drawRightString(8.02*inch,.22*inch,f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    out.seek(0)
    return out


def render_docx_v2(data: dict[str, Any], *, logo_stream: BytesIO | None = None) -> BytesIO:
    data = _canonicalize_input(data)
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.45);sec.bottom_margin=Inches(.58);sec.left_margin=Inches(.55);sec.right_margin=Inches(.55)
    sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)

    def rgb(hexv):
        v=str(hexv).replace("#","")
        return RGBColor(int(v[0:2],16),int(v[2:4],16),int(v[4:6],16))

    styles=doc.styles
    normal=styles["Normal"]
    normal.font.name="Aptos";normal.font.size=Pt(10.5);normal.font.color.rgb=rgb(INK)
    normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.08
    for name,size,color in [("Title",22,NAVY),("Heading 1",15,NAVY),("Heading 2",12,PRIMARY)]:
        st=styles[name];st.font.name="Aptos Display" if name!="Normal" else "Aptos";st.font.size=Pt(size);st.font.color.rgb=rgb(color)
        st.font.bold=True if name=="Title" else False
        st.paragraph_format.space_before=Pt(8 if name!="Title" else 0);st.paragraph_format.space_after=Pt(4)
        st.paragraph_format.keep_with_next=True

    def shade(cell, fill):
        tcPr=cell._tc.get_or_add_tcPr()
        shd=tcPr.find(qn("w:shd"))
        if shd is None:
            shd=OxmlElement("w:shd");tcPr.append(shd)
        shd.set(qn("w:fill"),fill.replace("#","").upper())

    def borders(cell, color=LINE, size="4"):
        tcPr=cell._tc.get_or_add_tcPr()
        borders_el=tcPr.first_child_found_in("w:tcBorders")
        if borders_el is None:
            borders_el=OxmlElement("w:tcBorders");tcPr.append(borders_el)
        for edge in ("top","left","bottom","right","insideH","insideV"):
            tag="w:"+edge
            el=borders_el.find(qn(tag))
            if el is None:
                el=OxmlElement(tag);borders_el.append(el)
            el.set(qn("w:val"),"single");el.set(qn("w:sz"),size);el.set(qn("w:color"),color.replace("#",""))

    def cell_text(cell,text,*,bold=False,size=9.5,color=INK,align=None):
        cell.text=""
        p=cell.paragraphs[0]
        p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
        if align is not None:p.alignment=align
        r=p.add_run(_txt(text));r.bold=bold;r.font.name="Aptos";r.font.size=Pt(size);r.font.color.rgb=rgb(color)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def add_table(headers,rows,widths=None,small=False):
        cols = len(headers) if headers else (len(rows[0]) if rows else 1)
        t=doc.add_table(rows=1 if headers else 0,cols=cols)
        t.autofit=True
        for i,h in enumerate(headers):
            cell_text(t.cell(0,i),h.upper(),bold=True,size=8.3,color=MUTED)
            shade(t.cell(0,i),"F2F5F7");borders(t.cell(0,i))
        for row in rows:
            cells=t.add_row().cells
            for i,val in enumerate(row):
                cell_text(cells[i],val,size=8.6 if small else 9.2)
                borders(cells[i])
        for row in t.rows:
            trPr=row._tr.get_or_add_trPr()
            cant=OxmlElement("w:cantSplit");trPr.append(cant)
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
        return t

    def heading(text,level=1,eyebrow=None):
        if eyebrow:
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.keep_with_next=True
            r=p.add_run(eyebrow.upper());r.bold=True;r.font.size=Pt(8.5);r.font.color.rgb=rgb(PRIMARY)
        h=doc.add_heading(text,level=level)
        h.paragraph_format.keep_with_next=True

    def kpis(items,cols=5):
        items=list(items)
        rows=((len(items)+cols-1)//cols)*2
        t=doc.add_table(rows=rows,cols=cols)
        idx=0
        for block in range(0,rows,2):
            for col in range(cols):
                if idx<len(items):
                    label,val,tone=items[idx]
                else:
                    label,val,tone="","",""
                cell_text(t.cell(block,col),label.upper(),bold=True,size=7.7,color=MUTED,align=WD_ALIGN_PARAGRAPH.CENTER)
                cell_text(t.cell(block+1,col),val,bold=True,size=10.2,color=tone or NAVY,align=WD_ALIGN_PARAGRAPH.CENTER)
                shade(t.cell(block,col),"F2F5F7");shade(t.cell(block+1,col),"FFFFFF")
                borders(t.cell(block,col));borders(t.cell(block+1,col))
                idx+=1
        doc.add_paragraph().paragraph_format.space_after=Pt(1)

    def two_panel(left_title,left_text,right_title,right_text,left_fill="F7FAFC",right_fill="F7FAFC"):
        t=doc.add_table(rows=2,cols=2)
        for i,(title,fill) in enumerate(((left_title,left_fill),(right_title,right_fill))):
            cell_text(t.cell(0,i),title.upper(),bold=True,size=8.5,color=NAVY);shade(t.cell(0,i),fill);borders(t.cell(0,i))
        cell_text(t.cell(1,0),left_text,size=9.5);cell_text(t.cell(1,1),right_text,size=9.5)
        borders(t.cell(1,0));borders(t.cell(1,1))
        doc.add_paragraph().paragraph_format.space_after=Pt(1)

    def add_chart(charts,key,width=6.9):
        stream=charts.get(key)
        if not stream:return
        stream.seek(0)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(stream,width=Inches(width))
        p.paragraph_format.space_after=Pt(4)

    brand=data.get("branding") or {};ident=data.get("identity") or {};valuation=data.get("valuation") or {};lenses=data.get("decision_lenses") or {};evidence=data.get("evidence") or {}
    charts=chart_bundle(data)

    # Header/footer.
    hp=sec.header.paragraphs[0]
    hp.text=(brand.get("title") or "Market Forensics").upper()+"  |  RESEARCH"
    hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:r.font.size=Pt(8);r.font.color.rgb=rgb(MUTED)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=fp.add_run(f"{brand.get('footer') or 'Lose Money Rules'} | Market Forensics | {ident.get('ticker','')} | ")
    r.font.size=Pt(8);r.font.color.rgb=rgb(MUTED)
    fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");fp._p.append(fld)

    if logo_stream is not None:
        try:
            logo_stream.seek(0);doc.add_picture(logo_stream,width=Inches(.34))
        except Exception:
            pass
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1)
    r=p.add_run((brand.get("title") or "Market Forensics").upper());r.bold=True;r.font.size=Pt(9);r.font.color.rgb=rgb(PRIMARY)
    title=doc.add_paragraph(style="Title");title.add_run(f"{ident.get('ticker','')}  |  {ident.get('company','')}")
    meta="  |  ".join(x for x in [ident.get("sector"),ident.get("industry"),ident.get("market_provider"),ident.get("market_as_of")] if x)
    p=doc.add_paragraph(meta or "Private CONTROL research");p.paragraph_format.space_after=Pt(6)
    for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=rgb(MUTED)

    # Hero.
    t=doc.add_table(rows=1,cols=2)
    left=t.cell(0,0);right=t.cell(0,1)
    left.text="";right.text=""
    p=left.paragraphs[0];r=p.add_run("RESEARCH CONCLUSION\n");r.bold=True;r.font.size=Pt(8);r.font.color.rgb=rgb(MUTED)
    r=p.add_run(_txt(data.get("conclusion")));r.bold=True;r.font.size=Pt(19);r.font.color.rgb=rgb(_tone(str(data.get("conclusion"))))
    p.add_run("\n"+f"Value {_txt(lenses.get('value'))} | Path {_txt(lenses.get('path'))} | Confidence {_txt(lenses.get('model_confidence'))}").font.size=Pt(8.8)
    p=right.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run("BASE TARGET\n");r.bold=True;r.font.size=Pt(8);r.font.color.rgb=rgb(MUTED)
    r=p.add_run(_money(valuation.get("base")));r.bold=True;r.font.size=Pt(20);r.font.color.rgb=rgb(PRIMARY)
    p.add_run("\n"+_pct(valuation.get("base_gap_pct"))+" vs market | "+_txt(valuation.get("base_quality")).replace("_"," ")).font.size=Pt(8.8)
    shade(left,"F7FAFC");shade(right,"F7FAFC");borders(left,"BDD0DF","6");borders(right,"BDD0DF","6")
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

    # Keep the Bear / Base / Bull visual directly below the conclusion hero in Word too.
    add_chart(charts,"valuation_map",6.9)
    kpis([
        ("Current price",_money(valuation.get("current_price")),MARKET),("Bear",_money(valuation.get("bear")),NEGATIVE),
        ("Base",_money(valuation.get("base")),PRIMARY),("Bull",_money(valuation.get("bull")),POSITIVE),("Base gap",_pct(valuation.get("base_gap_pct")),PRIMARY),
        ("Valuation quality",_txt(valuation.get("base_quality")).replace("_"," "),_tone(str(valuation.get("base_quality")))),
        ("Value lens",_txt(lenses.get("value")),_tone(str(lenses.get("value")))),("Expectations",_txt(lenses.get("expectations")),_tone(str(lenses.get("expectations")))),
        ("Variant",_txt(lenses.get("variant")),_tone(str(lenses.get("variant")))),("Path",_txt(lenses.get("path")),_tone(str(lenses.get("path")))),
        ("Model confidence",_txt(lenses.get("model_confidence")),_tone(str(lenses.get("model_confidence")))),("Thesis control",_txt(lenses.get("thesis_control")),_tone(str(lenses.get("thesis_control")))),
    ],6)

    thesis=data.get("thesis") or {}
    two_panel("MARKET VIEW",_txt(thesis.get("market_view"),500),"OUR VIEW / VARIANT",_txt(thesis.get("our_view"),500))
    must="\n".join("- "+_txt(x.get("text"),180) for x in (thesis.get("what_must_be_true") or [])[:3]) or "-"
    wrong="\n".join("- "+_txt(x.get("text"),180) for x in (thesis.get("what_proves_wrong") or [])[:3]) or "-"
    two_panel("WHAT MUST BE TRUE",must,"WHAT WOULD PROVE US WRONG",wrong)

    fort="\n".join("- "+_txt(r.get("label"))+": "+_txt(r.get("detail"),180) for r in (evidence.get("for") or [])[:4]) or "-"
    against="\n".join("- "+_txt(r.get("label"))+": "+_txt(r.get("detail"),180) for r in (evidence.get("against") or [])[:4]) or "-"
    two_panel("FOR",fort,"AGAINST",against,"EDF7F1","FBEFEF")

    company_quality=data.get("company_quality") or valuation.get("company_quality") or {}
    if company_quality:
        alarms=list(company_quality.get("alarm_bells") or [])
        strengths=list(company_quality.get("strengths") or [])
        add_table([],[
            ["Company quality",_txt(company_quality.get("state"))+" | "+_txt(company_quality.get("headline"),300)],
            ["Alarm bells"," | ".join(_txt(x.get("detail"),160) for x in alarms[:2]) or "No material automatic alarm bell."],
            ["Strengths"," | ".join(_txt(x.get("detail"),160) for x in strengths[:2]) or "No automatic strength clears the threshold yet."],
        ])

    tape=data.get("tape") or {};tm=tape.get("metrics") or {};cats=data.get("catalysts") or [];next_cat=cats[0] if cats else {};mon=data.get("monitoring") or {}
    add_table([],[
        ["Tape",f"Regime {_txt(tm.get('regime'))} | Rank {_txt(tm.get('rank'))} | Confidence {_txt(tm.get('confidence'))} | Net Tape {_num(tm.get('net_tape'),0)}"],
        ["Next catalyst"," | ".join(filter(None,[_txt(next_cat.get("event")),_txt(next_cat.get("timing")),_txt(next_cat.get("status"))]))],
        ["Invalidation",_txt(mon.get("thesis_invalidation"),320)],
    ])
    if evidence.get("score") is not None:
        p=doc.add_paragraph(f"Evidence diagnostic {float(evidence['score']):+.2f} - secondary diagnostic only; Decision Lenses remain canonical.")
        for r in p.runs:r.font.size=Pt(8.5);r.font.color.rgb=rgb(MUTED)

    if data.get("mode")=="executive":
        out=BytesIO();doc.save(out);out.seek(0);return out

    doc.add_page_break()
    heading("Valuation",1,"Decision economics")
    p=doc.add_paragraph("PRICE -> BEAR / BASE / BULL -> GAP. Method detail follows after the decision range.")
    for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=rgb(MUTED)
    kpis([("Current price",_money(valuation.get("current_price")),MARKET),("Bear",_money(valuation.get("bear")),NEGATIVE),("Base",_money(valuation.get("base")),PRIMARY),("Bull",_money(valuation.get("bull")),POSITIVE),("Expected value",_money(valuation.get("expected_value")),NAVY)],5)
    add_chart(charts,"price_context",6.9)
    rows=[]
    for row in valuation.get("scenarios") or []:
        methods={m.get("key"):m for m in row.get("methods") or []}
        rows.append([row.get("name"),_pct((_n(row.get("probability")) or 0)*100,signed=False),_money(row.get("target")),_txt(row.get("quality")).replace("_"," "),_money((methods.get("pe") or {}).get("value")),_money((methods.get("ev_sales") or {}).get("value")),_money((methods.get("fcf_yield") or {}).get("value")),_money(row.get("dcf"))])
    add_table(["Case","Probability","Target","Quality","P/E","EV/Sales","FCF Yield","DCF"],rows,small=True)
    weight_rows=[]
    for row in valuation.get("scenarios") or []:
        methods={m.get("key"):m for m in row.get("methods") or []}
        def method_weight(key):
            weight=_n((methods.get(key) or {}).get("weight"))
            return _pct(weight*100,signed=False) if weight is not None else "-"
        weight_rows.append([row.get("name"),method_weight("pe"),method_weight("ev_sales"),method_weight("fcf_yield")])
    if weight_rows:
        p=doc.add_paragraph();r=p.add_run("EFFECTIVE METHOD WEIGHTS");r.bold=True;r.font.size=Pt(8.5);r.font.color.rgb=rgb(PRIMARY)
        add_table(["Case","P/E weight","EV/Sales weight","FCF Yield weight"],weight_rows,small=True)
    share=valuation.get("share_basis") or {}
    p=doc.add_paragraph("Share denominator: "+_num(share.get("shares"),1)+" | Source: "+_txt(share.get("source")).replace("_"," ")+" | Verified: "+("YES" if share.get("verified") else "NO"))
    for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=rgb(MUTED)
    if valuation.get("provisional"):
        p=doc.add_paragraph();r=p.add_run("PROVISIONAL VALUATION. ");r.bold=True;r.font.color.rgb=rgb(CAUTION);p.add_run("Stored targets remain visible but are not decision-grade intrinsic evidence until the quality issue is resolved.")
    for warning in (valuation.get("warnings") or [])[:4]:
        p=doc.add_paragraph("WATCH - "+_txt(warning,420));p.style=styles["Normal"]

    heading("Quality → valuation",1,"Explicit automatic price adjustments")
    policy=valuation.get("valuation_policy") or {};ledger=valuation.get("valuation_impact_ledger") or []
    kpis([
        ("Company quality",_txt((valuation.get("company_quality") or {}).get("state") or "INSUFFICIENT EVIDENCE"),NAVY),
        ("Risk premium","+"+str(int(_n(policy.get("risk_premium_bps")) or 0))+" bps",NAVY),
        ("Growth haircut","-"+str(int(_n(policy.get("growth_haircut_bps")) or 0))+" bps",NAVY),
        ("Terminal haircut","-"+str(int(_n(policy.get("terminal_growth_haircut_bps")) or 0))+" bps",NAVY),
        ("Bear probability","+"+str(int(_n(policy.get("bear_probability_shift_pts")) or 0))+" pts",NAVY),
    ],5)
    if ledger:
        add_table(["Item","Effect","Impact","Reason"],[[r.get("item"),_txt(r.get("effect")).replace("_"," "),r.get("impact"),_txt(r.get("reason"),300)] for r in ledger[:10]],small=True)
    p=doc.add_paragraph("Positive company quality receives no automatic premium. Observed operating strength already enters through growth, margins, returns and cash flows.")
    for r in p.runs:r.font.size=Pt(8.5);r.font.color.rgb=rgb(MUTED)

    heading("Thesis / variant",1,"Why the market may be wrong")
    two_panel("MARKET VIEW",_txt(thesis.get("market_view"),800),"OUR VIEW / VARIANT",_txt(thesis.get("our_view"),800))
    two_panel("WHAT MUST BE TRUE",must,"WHAT WOULD PROVE US WRONG",wrong)
    if thesis.get("variant_evidence"):
        p=doc.add_paragraph();r=p.add_run("VARIANT EVIDENCE - ");r.bold=True;p.add_run(_txt(thesis.get("variant_evidence"),1000))

    heading("Expectations",1,"What is priced")
    implied=(data.get("expectations") or {}).get("implied") or {}
    if implied.get("available"):
        rows=[]
        for row in implied.get("drivers") or []:
            unit=row.get("unit");mv=_n(row.get("market_implied"));bv=_n(row.get("base"))
            rows.append([row.get("label"),_pct(mv*100,signed=False) if unit=="%" and mv is not None else _num(mv,1,"x"),_pct(bv*100,signed=False) if unit=="%" and bv is not None else _num(bv,1,"x"),row.get("read")])
        add_table(["Driver","Market-implied","Our Base","Read"],rows)
    exp=(data.get("expectations") or {}).get("rows") or []
    if exp:
        add_table(["Metric","Period","Market","Our view","Confidence"],[[r.get("metric"),r.get("period"),str(r.get("market") if r.get("market") is not None else "-"),str(r.get("ours") if r.get("ours") is not None else "-"),r.get("confidence")] for r in exp[:12]])

    heading("Evidence for / against",1,"Research evidence")
    two_panel("FOR",fort,"AGAINST",against,"EDF7F1","FBEFEF")

    heading("Company quality",1,"Is the economic engine actually good?")
    cq=data.get("company_quality") or {}
    if cq:
        p=doc.add_paragraph();r=p.add_run(_txt(cq.get("state"))+" · ");r.bold=True;p.add_run(_txt(cq.get("headline"),600))
        add_table(["Dimension","State","Read"],[[d.get("label"),_txt(d.get("state")).replace("_"," "),_txt(d.get("detail"),320)] for d in (cq.get("dimensions") or [])],small=True)
        for flag in (cq.get("alarm_bells") or [])[:6]:
            p=doc.add_paragraph();r=p.add_run("ALARM · "+_txt(flag.get("severity"))+" · ");r.bold=True;p.add_run(_txt(flag.get("detail"),420))
    else:
        doc.add_paragraph("Company-quality evidence is not materialized yet.")

    heading("Business",1,"Operating reality")
    doc.add_paragraph(_txt((data.get("business") or {}).get("summary") or "Business research is not yet documented."))

    heading("Fundamentals",1,"Trend before table")
    ff=(data.get("fundamentals") or {}).get("forensics") or {}
    if ff:
        p=doc.add_paragraph();r=p.add_run(_txt(ff.get("state") or "DATA REVIEW")+" · ");r.bold=True;p.add_run(_txt(ff.get("headline"),600))
        strengths=list(ff.get("strengths") or [])[:4]
        risks=(list(ff.get("red_flags") or [])+list(ff.get("watches") or []))[:5]
        two_panel(
            "STRENGTHS",
            "\n".join("- "+_txt(x.get("label"),80)+": "+_txt(x.get("detail"),190) for x in strengths) or "-",
            "RED FLAGS / WATCH",
            "\n".join("- "+_txt(x.get("label"),80)+": "+_txt(x.get("detail"),190) for x in risks) or "-",
            "EDF7F1","FBEFEF",
        )
        for item in (list(ff.get("inconsistencies") or [])+list(ff.get("data_gaps") or []))[:5]:
            p=doc.add_paragraph();r=p.add_run("REVIEW · "+_txt(item.get("label"),90)+" · ");r.bold=True;p.add_run(_txt(item.get("detail"),380))
    add_chart(charts,"revenue_profitability",6.9);add_chart(charts,"cash_conversion",6.9);add_chart(charts,"working_capital",6.9)
    current=(data.get("fundamentals") or {}).get("current") or {}
    kpis([
        ("Basis",_txt(current.get("period") or current.get("comparison_basis")),NAVY),("Op margin",_pct(current.get("operating_margin_pct"),signed=False),NAVY),
        ("FCF margin",_pct(current.get("fcf_margin_pct"),signed=False),NAVY),("CFO / NI",_num(current.get("cfo_to_net_income"),2,"x"),NAVY),("ROIC",_pct(current.get("roic_pct"),signed=False),NAVY),
        ("Net debt / FCF",_num(current.get("net_debt_to_fcf"),1,"x"),NAVY),("Inventory / Rev",_pct(current.get("inventory_to_revenue_pct"),signed=False),NAVY),
        ("Receivables / Rev",_pct(current.get("receivables_to_revenue_pct"),signed=False),NAVY),("CCC",_num(current.get("ccc"),0,"d"),NAVY),("Share count growth",_pct(current.get("share_count_growth_pct")),NAVY),
    ],5)
    if current.get("economic_reality_quality"):
        heading("Economic Reality",2,"Reported accounting → economic interpretation")
        kpis([
            ("Econ quality",_txt(current.get("economic_reality_quality")),NAVY),
            ("Reported net debt",_money(current.get("reported_net_debt")),NAVY),
            ("Economic net debt",_money(current.get("economic_net_debt")),NAVY),
            ("Operating leases",_money(current.get("operating_lease_liability")),NAVY),
            ("Lease / liabilities",_pct(current.get("operating_lease_share_of_liabilities_pct"),signed=False),NAVY),
            ("Revenue / leases",_num(current.get("lease_revenue_productivity_x"),2,"x"),NAVY),
            ("Growth capex proxy",_money(current.get("growth_capex_proxy")),NAVY),
            ("Owner-cash proxy",_money(current.get("owner_cash_proxy")),NAVY),
            ("FCF after SBC",_money(current.get("fcf_after_sbc")),NAVY),
            ("Lease-adj ROIC",_pct(current.get("lease_adjusted_roic_pct"),signed=False),NAVY),
        ],5)
        for flag in (current.get("economic_reality_flags") or [])[:6]:
            p=doc.add_paragraph()
            p.add_run(_txt(flag.get("code")).replace("_"," ")+" · ").bold=True
            p.add_run(_txt(flag.get("detail"),360))
        if current.get("economic_reality_unresolved"):
            p=doc.add_paragraph("ACCOUNTING REVIEW · material classification is unresolved; directional scoring is conservative until verified.")
            for r in p.runs:r.bold=True
    hist=(data.get("fundamentals") or {}).get("history") or []
    if hist:
        add_table(["Period","Rev growth","Gross M","Op M","FCF M","CFO/NI","ROIC","Inv/Rev","Rec/Rev","CCC"],[[r.get("period"),_pct(r.get("revenue_growth_pct")),_pct(r.get("gross_margin_pct"),signed=False),_pct(r.get("operating_margin_pct"),signed=False),_pct(r.get("fcf_margin_pct"),signed=False),_num(r.get("cfo_to_net_income"),2,"x"),_pct(r.get("roic_pct"),signed=False),_pct(r.get("inventory_to_revenue_pct"),signed=False),_pct(r.get("receivables_to_revenue_pct"),signed=False),_num(r.get("ccc"),0,"d")] for r in hist[-8:]],small=True)

    heading("Financial flows",1,"Follow the money")
    steps=_flow_rows(data)
    if steps:
        add_table(
            ["Flow","Line item","Signed amount","Route","Basis"],
            [[r.get("flow"),r.get("label"),_money(r.get("value")),r.get("route"),r.get("basis")] for r in steps[:16]],
            small=True,
        )
        p=doc.add_paragraph("Signed negatives are retained exactly as stored. The report does not invent balancing values to force a bridge.")
        for r in p.runs:r.font.size=Pt(8.5);r.font.color.rgb=rgb(MUTED)
    else:
        doc.add_paragraph(_txt((data.get("flows") or {}).get("summary") or "Materialized financial-flow bridge unavailable."))

    doc.add_page_break()
    heading("Management",1,"Execution accountability")
    mg=data.get("management") or {};eng=mg.get("engine") or {}
    kpis([("Execution",_txt(eng.get("label") or eng.get("execution_label") or "LOW DATA"),_tone(str(eng.get("label")))),("Evidence coverage",_pct(eng.get("coverage_pct"),signed=False),NAVY),("Promises",str(len(mg.get("promises") or [])),NAVY),("Scorable",str(eng.get("scorable") or "-"),NAVY),("Missed",str(eng.get("missed") or "-"),NEGATIVE)],5)
    promises=mg.get("promises") or []
    if promises:
        prows=[]
        for row in promises[:16]:
            lo=row.get("low");hi=row.get("high");unit=row.get("unit") or "";promise=_txt(row.get("target_text"))
            if promise=="-" and lo is not None:
                promise=str(lo) if lo==hi else f"{lo} - {hi}"
                if unit:promise+=" "+str(unit)
            prows.append([row.get("target_period") or (("FY"+str(row.get("target_year"))) if row.get("target_year") else "UNRESOLVED"),row.get("metric"),promise,str(row.get("actual") if row.get("actual") is not None else "-"),row.get("status"),row.get("comparability") or row.get("comparability_reason") or "-"])
        add_table(["Period","Metric","Promise","Actual","Status","Comparability"],prows,small=True)
    elif mg.get("summary"):
        doc.add_paragraph(_txt(mg.get("summary")))

    heading("Catalysts / bear case",1,"Path and failure modes")
    if cats:
        add_table(["Event","Timing","Direction","Status"],[[r.get("event"),r.get("timing"),r.get("direction"),r.get("status")] for r in cats[:12]])
    bears=data.get("bear_case") or []
    if bears:
        add_table(["Risk","Severity","Status","Invalidates","Why it matters"],[[r.get("risk"),r.get("severity"),r.get("status"),"YES" if r.get("invalidates") else "NO",_txt(r.get("evidence"),220)] for r in bears[:12]],small=True)

    heading("Tape & positioning",1,"Market evidence")
    kpis([("Regime",_txt(tm.get("regime") or tm.get("forensic_regime")),_tone(str(tm.get("regime")))),("Rank",_txt(tm.get("rank")),NAVY),("Confidence",_txt(tm.get("confidence")),NAVY),("Large flow",_money(tm.get("net_large")),NAVY),("Whale flow",_money(tm.get("net_whale")),NAVY),("Short pressure",_num(tm.get("bear_pressure"),0),NEGATIVE),("Absorption",_num(tm.get("absorption"),0),POSITIVE),("Net Tape",_num(tm.get("net_tape"),0),PRIMARY),("What changed",_txt(tape.get("what_changed"),60),NAVY),("Regime change",_txt(tape.get("what_would_change_regime"),60),NAVY)],5)
    add_chart(charts,"tape_price_flow",6.9);add_chart(charts,"tape_pressure",6.9)

    heading("Thesis invalidation / monitoring",1,"Decision control")
    p=doc.add_paragraph();r=p.add_run("THESIS INVALIDATION - ");r.bold=True;p.add_run(_txt(mon.get("thesis_invalidation"),900))
    if mon.get("locked_at"):
        p=doc.add_paragraph("Locked pre-investment basis: "+str(mon.get("locked_at"))+". Thresholds are not rewritten retroactively.")
        for r in p.runs:r.font.size=Pt(8.5);r.font.color.rgb=rgb(MUTED)
    rules=mon.get("rules") or []
    if rules:
        mrows=[]
        for row in rules[:20]:
            threshold=_txt(row.get("threshold_text"))
            if threshold=="-" and row.get("threshold") is not None:
                threshold=_num(row.get("threshold"),2)+((" "+row.get("unit")) if row.get("unit") else "")
            mrows.append([row.get("name") or row.get("metric"),threshold,row.get("operator"),_num(row.get("current_value"),2)+((" "+row.get("unit")) if row.get("unit") and row.get("current_value") is not None else ""),row.get("status"),str(row.get("last_observation") or "")[:19],"YES" if row.get("locked_pre_investment") else "NO"])
        add_table(["Metric / rule","Threshold","Direction","Current","State","Last observation","Locked"],mrows,small=True)

    heading("Validation",1,"Historical calibration")
    val=data.get("validation") or {}
    if not val.get("available"):
        p=doc.add_paragraph();r=p.add_run("NOT RUN. ");r.bold=True;p.add_run("No historical validation result is stored. No validation confidence is implied.")
    else:
        kpis([("State",_txt(val.get("state")),_tone(str(val.get("state")))),("Reliability",_pct(val.get("reliability"),signed=False),NAVY),("Samples",str(val.get("sample_count") or 0),NAVY),("Valuation accuracy",_pct(val.get("valuation_accuracy"),signed=False),NAVY),("Direction accuracy",_pct(val.get("direction_accuracy"),signed=False),NAVY),("Range coverage",_pct(val.get("range_coverage"),signed=False),NAVY),("Assumption accuracy",_pct(val.get("assumption_accuracy"),signed=False),NAVY),("Requested span",str(val.get("lookback_years") or "-")+"Y",NAVY),("History span",_txt(val.get("history_span")),NAVY),("Status",_txt(val.get("status")),NAVY)],5)
        add_chart(charts,"validation",6.9)

    heading("Sources / audit",1,"Provenance")
    sources=data.get("sources") or []
    if sources:
        source_table=add_table(["Provider","Document / type","Published","Retrieved","Title"],[[r.get("provider")," / ".join(x for x in [r.get("document"),r.get("type")] if x),str(r.get("published_at") or "")[:10],str(r.get("retrieved_at") or "")[:19],_txt(r.get("title"),160)] for r in sources[:50]],small=True)
        for source_row in source_table.rows:
            for cell in source_row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size=Pt(6.7);run.font.color.rgb=rgb(MUTED)
    contract=data.get("data_contract") or {}
    p=doc.add_paragraph(f"Report data contract {data.get('contract_version','')} | Cache {contract.get('materialized_cache_event') or '-'} | Provider refresh started: {'YES' if contract.get('provider_refresh_started') else 'NO'} | Heavy analytics started: {'YES' if contract.get('heavy_analytics_started') else 'NO'}")
    for r in p.runs:r.font.size=Pt(8);r.font.color.rgb=rgb(MUTED)

    out=BytesIO();doc.save(out);out.seek(0);return out


__all__=["render_pdf_v2","render_docx_v2"]
